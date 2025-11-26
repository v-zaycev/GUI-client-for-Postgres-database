from PyQt6.QtWidgets import ( QVBoxLayout, QPushButton, QLabel, QHBoxLayout, QToolButton, QMenu)
from PyQt6.QtGui import QAction
from sources.base_client import PgsqlClient
from sources.main_window.basic_widget import BasicWidget

import pandas as pd
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import getSampleStyleSheet
import os

from sources.main_window.names_conversion import names_conversion
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class ReportsWidget(BasicWidget):
    def __init__(self, db_client : PgsqlClient):
        super().__init__()
        self.pgsql_client = db_client
        layout = QVBoxLayout()
        self.setLayout(layout)
        # Заголовок
        self.title_label = QLabel("Отчёты")
        self.title_label.setStyleSheet("font-size: 16px; font-weight: bold; margin: 10px;")
        layout.addWidget(self.title_label)
        
        # Панель кнопок
        self.buttons_layout = self.init_buttons()
        self.layout().addLayout(self.buttons_layout)
        self.load_button.clicked.connect(self.load_data)
        self.export_xlsx.clicked.connect(self.create_xlsx_report)
        self.export_csv.clicked.connect(self.create_csv_report)
        self.export_pdf.clicked.connect(self.export_to_pdf)

        self.tool_btn = QToolButton()
        self.tool_btn.setText("Выбор отчёта")  # Текст кнопки
        self.tool_btn.setPopupMode(QToolButton.ToolButtonPopupMode.MenuButtonPopup)

        # Создаем выпадающее меню
        self.dropdown_menu = QMenu()
        action1 = QAction("Статистика по палатам", self)
        action2 = QAction("Статистика по диагнозам", self)
        self.dropdown_menu.addAction(action1)
        self.dropdown_menu.addAction(action2)

        # Привязываем меню к кнопке
        self.tool_btn.setMenu(self.dropdown_menu)

        # Подключаем обработчики
        action1.triggered.connect(lambda: self.set_table("ward_occupancy_report"))
        action2.triggered.connect(lambda: self.set_table("diagnosis_statistics"))
        layout.addWidget(self.tool_btn)

        # Таблица для отображения данных
        self.table = self.init_table()
        layout.addWidget(self.table)

        # Статус
        self.status_label = QLabel("Нажмите кнопку для загрузки данных")
        layout.addWidget(self.status_label)

        self.rights = {"ward_occupancy_report" : self.pgsql_client.get_table_rights("ward_occupancy_report"),
                       "diagnosis_statistics" : self.pgsql_client.get_table_rights("diagnosis_statistics")}
        self.set_table("ward_occupancy_report")

    def set_table(self, table : str):
        self.current_table_name = table
        self.data, self.description = self.pgsql_client.select(["*"], table)
        self.show_data(self.data, self.description)
        if self.current_table_name == 'ward_occupancy_report':
            self.title_label.setText("Статистика по палатам")
        elif self.current_table_name == 'diagnosis_statistics':
            self.title_label.setText("Статистика по диагнозам")
    
    def create_csv_report(self):
        reports_dir = "reports"
        filepath = os.path.join(reports_dir, f"{self.current_table_name}.csv")

        df = pd.DataFrame(data = self.data, columns=[names_conversion[self.current_table_name][desc[0]] for desc in self.description])
        df.to_csv(filepath)

    def init_buttons(self) -> QHBoxLayout:
        buttons_layout = QHBoxLayout()

        self.load_button = QPushButton("Обновить данные")
        self.load_button.setEnabled(True)
        self.load_button.setFixedWidth(120)
        buttons_layout.addWidget(self.load_button)
        
        # Кнопки
        self.export_xlsx = QPushButton("Экспортировать в .xlsx")
        self.export_csv = QPushButton("Экспортировать в .csv") 
        self.export_pdf = QPushButton("Экспортировать в .pdf")
        
        self.export_xlsx.setEnabled(True)
        self.export_csv.setEnabled(True)
        self.export_pdf.setEnabled(True)
        
        # Фиксируем ширину для аккуратного вида
        self.export_xlsx.setFixedWidth(150)
        self.export_csv.setFixedWidth(150)
        self.export_pdf.setFixedWidth(150)
        
        buttons_layout.addWidget(self.export_xlsx)
        buttons_layout.addWidget(self.export_csv)
        buttons_layout.addWidget(self.export_pdf)
        buttons_layout.addStretch()
                
        return buttons_layout
    
    def create_xlsx_report(self):
        reports_dir = "reports"
        filepath = os.path.join(reports_dir, f"{self.current_table_name}.xlsx")
    
        df = pd.DataFrame(data = self.data, columns=[names_conversion[self.current_table_name][desc[0]] for desc in self.description])
        df.to_excel(filepath, index=False, engine='openpyxl')
        
        workbook = load_workbook(filepath)
        worksheet = workbook.active
        
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        workbook.save(filepath)

    def export_to_pdf(self):

        try:
            font_paths = ["C:/Windows/Fonts/arial.ttf",
                          "C:/Windows/Fonts/times.ttf"]
            
            font_registered = False
            for font_path in font_paths:
                if os.path.exists(font_path):
                    pdfmetrics.registerFont(TTFont('CyrillicFont', font_path))
                    font_registered = True
                    break
                    
            if not font_registered:
                pdfmetrics.registerFont(TTFont('CyrillicFont', 'Helvetica'))
        except:
            pdfmetrics.registerFont(TTFont('CyrillicFont', 'Helvetica'))


        reports_dir = "reports"
        filepath = os.path.join(reports_dir, f"{self.current_table_name}.pdf")
        doc = SimpleDocTemplate(filepath, pagesize=landscape(A4))
        elements = []
        
        df = pd.DataFrame(data = self.data, columns=[names_conversion[self.current_table_name][desc[0]] for desc in self.description])
        data = [df.columns.tolist()] + df.values.tolist()
        
        styles = getSampleStyleSheet()
        title_style = styles['Title']
        title_style.fontName = 'CyrillicFont'
    
        title_paragraph = Paragraph("Статистика по палатам" if self.current_table_name == 'ward_occupancy_report' else "Статистика по диагнозам"
                                    , title_style)
        elements.append(title_paragraph)
        elements.append(Spacer(1, 20)) 

        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, -1), 'CyrillicFont'),
            ('FONTSIZE', (0, 0), (-1, 0), 14),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        doc.build(elements)
        